from pathlib import Path
from util import normalize
from openpyxl import load_workbook
from collections import defaultdict
import os


def filter_duplicates_and_rename(excel_path, results):

    existing_keys, existing_names = load_existing_data(excel_path)

    incoming_keys = set()
    name_counts = defaultdict(int)

    # 既存ファイル名からカウント初期化
    for name in existing_names:
        base = Path(name).stem.split("_")[0]
        name_counts[base] += 1

    filtered = []
    duplicates = []

    for r in results:

        # 重複キー
        key = (r["date"], r["amount"], r["shop"] or "")

        # 重複判定s
        if key in existing_keys:
            duplicates.append((r, "既存Excelと重複"))
            continue

        if key in incoming_keys:
            duplicates.append((r, "レシートファイル内で重複"))
            continue

        incoming_keys.add(key)

        # ファイル名生成
        date = r["date"]
        ext = Path(r["image"]).suffix

        if not date:
            new_name = Path(r["image"]).name
        else:
            name_counts[date] += 1
            if name_counts[date] == 1:
                new_name = f"{date}{ext}"
            else:
                new_name = f"{date}_{name_counts[date]}{ext}"

        # rename
        old_path = Path(r["image"])
        new_path = old_path.parent / new_name

        # 念のため衝突回避
        counter = 1
        while new_path.exists():
            new_path = old_path.parent / f"{date}_{counter}{ext}"
            counter += 1

        os.rename(old_path, new_path)

        r["image"] = str(new_path)
        filtered.append(r)

    return filtered, duplicates


def load_existing_data(excel_path):

    keys = set()
    names = set()

    if not Path(excel_path).exists():
        return keys, names

    wb = load_workbook(excel_path)
    ws = wb.active

    for row in ws.iter_rows(min_row=2, values_only=True):

        filename = row[0]
        date = normalize(row[1])
        amount = normalize(row[2])
        shop = normalize(row[3] or "")

        if filename:
            names.add(filename)

        keys.add((date, amount, shop))

    return keys, names
