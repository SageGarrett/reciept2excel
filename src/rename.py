from collections import defaultdict
import os
from pathlib import Path
from openpyxl import load_workbook


def rename_files(excel_path, unique_ocr_results):

    # 既存ファイル名の取得
    existing_names = load_existing_filenames(excel_path)

    # 新しいファイル名の生成
    new_names = generate_new_filenames(unique_ocr_results, existing_names)

    # ファイルのリネームとresultsの更新
    apply_rename(unique_ocr_results, new_names)


def apply_rename(results, new_names):

    for r, new_name in zip(results, new_names):

        old_path = Path(r["image"])
        new_path = old_path.parent / new_name

        os.rename(old_path, new_path)

        # results更新
        r["image"] = str(new_path)


def generate_new_filenames(results, existing_names):

    # 日付ベースのファイル名カウント（重複があればインクリメントし、連番として付与する）
    name_counts = defaultdict(int)
    # 新しいファイル名のリスト
    new_names = []

    # 既存ファイルからカウント初期化
    for name in existing_names:
        file_name = name.split(".")[0].split("_")[0]
        name_counts[file_name] += 1

    # resultsに対して決定
    for r in results:

        date = r["date"]
        ext = Path(r["image"]).suffix

        # 日付がない場合は元のファイル名を使用
        if not date:
            original_name = Path(r["image"]).name
            new_names.append(original_name)
            continue

        name_counts[date] += 1

        if name_counts[date] == 1:
            new_name = f"{date}{ext}"
        else:
            new_name = f"{date}_{name_counts[date]}{ext}"

        new_names.append(new_name)

    return new_names


def load_existing_filenames(excel_path, filename_col=1):

    existing_names = set()

    if not Path(excel_path).exists():
        return existing_names

    try:
        wb = load_workbook(excel_path)
        ws = wb.active

        for row in ws.iter_rows(min_row=2, values_only=True):
            filename = row[filename_col - 1]
            if filename:
                existing_names.add(filename)

    except FileNotFoundError:
        pass

    return existing_names
