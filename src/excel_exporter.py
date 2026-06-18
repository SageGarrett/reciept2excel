import streamlit as st
from config import DEBUG_DIR
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from pathlib import Path
from datetime import datetime


def export_to_excel_from_results(results: list[dict], output_path: str):

    headers = ["ファイル名", "日付", "金額", "支払先", "警告"]

    # 初回（ファイルがない）
    if not Path(output_path).exists():

        wb = Workbook()
        ws = wb.active

        # ヘッダー作成
        ws.append(headers)

    else:
        wb = load_workbook(output_path)
        ws = wb.active

    yellow_fill = PatternFill(
        start_color="FFFF99", end_color="FFFF99", fill_type="solid"
    )

    # 列番号取得
    col_map = {cell.value: idx for idx, cell in enumerate(ws[1], start=1)}

    file_col = col_map.get("ファイル名")
    date_col = col_map.get("日付")
    amount_col = col_map.get("金額")
    shop_col = col_map.get("支払先")
    warning_col = col_map.get("警告")

    if None in (date_col, amount_col, shop_col, warning_col):
        raise ValueError("必要な列が見つかりません")

    # 追記開始行
    row = ws.max_row + 1

    # データ追加
    for r in results:

        # 値
        image_name = Path(r["image"]).name
        date_value = to_excel_date(r["date"])
        amount_value = r["amount"]
        shop_value = r["shop"]

        # セル
        file_cell = ws.cell(row=row, column=file_col)
        date_cell = ws.cell(row=row, column=date_col)
        date_cell.number_format = "yyyy-mm-dd"
        amount_cell = ws.cell(row=row, column=amount_col)
        amount_cell.number_format = '"¥"#,##0'
        shop_cell = ws.cell(row=row, column=shop_col)
        warning_cell = ws.cell(row=row, column=warning_col)

        # 書き込み
        file_cell.value = image_name
        date_cell.value = date_value
        amount_cell.value = amount_value
        shop_cell.value = shop_value

        # === デバッグ出力 ===
        debug_dir = DEBUG_DIR
        debug_dir.mkdir(exist_ok=True)

        debug_path = debug_dir / (image_name + ".txt")

        with open(debug_path, "w", encoding="utf-8") as f:
            f.write("=== STRUCTURED DATA ===\n")
            f.write(str(r) + "\n\n")

        row += 1

        warn_flag = False

        # 欠損チェック
        if not date_value:
            date_cell.fill = yellow_fill
            warn_flag = True

        if not amount_value:
            amount_cell.fill = yellow_fill
            warn_flag = True

        if not shop_value:
            shop_cell.fill = yellow_fill
            warn_flag = True

        # 決算年度チェック
        if date_value:
            try:
                if not is_in_fiscal_year(date_value):
                    date_cell.fill = yellow_fill
                    warn_flag = True
            except Exception as e:
                st.error(f"決算期判定エラー: {e}")
                warn_flag = True

        # 警告
        warning_cell.value = warn_flag

    # 列幅調整
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter

        for cell in col:
            if cell.value:

                if isinstance(cell.value, datetime):
                    text = cell.value.strftime("%Y-%m-%d")
                else:
                    text = str(cell.value)

                length = 0
                for c in text:
                    if ord(c) > 256:
                        length += 2
                    else:
                        length += 1

                max_length = max(max_length, length)

        ws.column_dimensions[col_letter].width = max_length + 3

    wb.save(output_path)


def to_excel_date(v):
    if v is None:
        return None

    try:
        return datetime.strptime(v, "%Y-%m-%d")
    except:
        return None


def is_in_fiscal_year(dt: datetime) -> bool:

    fiscal_year = st.session_state["fiscal_year"]
    fiscal_month = st.session_state["fiscal_month"]

    # 決算月の翌月が期首
    if fiscal_month == 12:
        start = datetime(fiscal_year, 1, 1)
        end = datetime(fiscal_year + 1, 1, 1)
    else:
        start = datetime(fiscal_year - 1, fiscal_month + 1, 1)
        end = datetime(fiscal_year, fiscal_month + 1, 1)

    return start <= dt < end
